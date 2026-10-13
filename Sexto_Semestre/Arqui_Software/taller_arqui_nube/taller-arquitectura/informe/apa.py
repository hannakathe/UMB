"""Utilidades de formato APA 7 (versión estudiante) sobre python-docx."""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FONT = "Times New Roman"
INDENT = Cm(1.27)
TEXT_WIDTH_IN = 6.5        # carta (8,5") - 2 × 2,54 cm
MAX_IMG_H_IN = 7.4


def _set_font(run_or_style, size=12, name=FONT):
    f = run_or_style.font
    f.name = name
    f.size = Pt(size)
    rpr = run_or_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), name)


def _field(paragraph, instr):
    """Inserta un campo de Word (PAGE, TOC...)."""
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "Clic derecho → Actualizar campo para generar la tabla de contenido." if "TOC" in instr else "1"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, t, s):
        run._r.append(el)
    r2 = paragraph.add_run(); r2._r.append(txt)
    r3 = paragraph.add_run(); r3._r.append(e)
    return run


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(table, spec):
    """spec: dict borde->(tamaño, color) para tblBorders (top, bottom, left, right, insideH, insideV)."""
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in spec:
            sz, color = spec[edge]
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tblPr.append(b)


def _cell_bottom_border(cell, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    bt = OxmlElement("w:bottom"); bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), str(sz)); bt.set(qn("w:color"), "000000")
    tb.append(bt); tcPr.append(tb)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)


def _keep_with_next(p):
    p.paragraph_format.keep_with_next = True


INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[\[[^\]]+\]\])")


class APA:
    def __init__(self, base: Path):
        self.base = base
        self.doc = Document()
        self.fig = 0
        self.tab = 0
        self.missing = []
        self._setup()

    # ------------------------------------------------------------ estilos
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, m, Cm(2.54))
        sec.header_distance = Cm(1.27)
        normal = d.styles["Normal"]
        _set_font(normal)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_before = Pt(0); pf.space_after = Pt(0)
        for lvl in (1, 2, 3):
            st = d.styles[f"Heading {lvl}"]
            _set_font(st)
            st.font.bold = True
            st.font.italic = (lvl == 3)
            st.font.color.rgb = RGBColor(0, 0, 0)
            f = st.paragraph_format
            f.space_before = Pt(0); f.space_after = Pt(0)
            f.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            f.first_line_indent = Cm(0)
            f.alignment = WD_ALIGN_PARAGRAPH.CENTER if lvl == 1 else WD_ALIGN_PARAGRAPH.LEFT
            f.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            st = d.styles[name]
            _set_font(st)
        for lvl in (1, 2, 3):
            try:
                st = d.styles[f"TOC {lvl}"]
            except KeyError:
                st = d.styles.add_style(f"TOC {lvl}", 1)
            _set_font(st)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
        # número de página arriba a la derecha (APA 7)
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = _field(hp, "PAGE")
        for run in hp.runs:
            _set_font(run)

    # ------------------------------------------------------------ texto
    def _runs(self, p, text, size=12, bold=False, italic=False, highlight_placeholders=True):
        for part in INLINE.split(text):
            if not part:
                continue
            b, i, mono, hl = bold, italic, False, False
            if part.startswith("**") and part.endswith("**"):
                self._runs(p, part[2:-2], size=size, bold=True, italic=italic)
                continue
            elif part.startswith("`") and part.endswith("`"):
                part, mono = part[1:-1], True
            elif part.startswith("[[") and part.endswith("]]"):
                part, hl = "[" + part[2:-2] + "]", True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                part, i = part[1:-1], True
            run = p.add_run(part)
            _set_font(run, size=size - (1 if mono else 0), name="Courier New" if mono else FONT)
            if mono and b:
                run.bold = True
            run.bold, run.italic = b, i
            if hl:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    def p(self, text, indent=True, align=None, size=12, bold=False, italic=False, keep=False):
        para = self.doc.add_paragraph()
        para.paragraph_format.first_line_indent = INDENT if indent else Cm(0)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._runs(para, text, size=size, bold=bold, italic=italic)
        if keep:
            _keep_with_next(para)
        return para

    def h1(self, t): return self._h(t, 1)
    def h2(self, t): return self._h(t, 2)
    def h3(self, t): return self._h(t, 3)

    def _h(self, t, lvl):
        h = self.doc.add_heading(level=lvl)
        self._runs(h, t, bold=True, italic=(lvl == 3))
        return h

    def bullets(self, items, numbered=False):
        for it in items:
            para = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            self._runs(para, it)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def code(self, text, size=8.5):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, {e: (4, "999999") for e in ("top", "bottom", "left", "right")})
        c = t.rows[0].cells[0]
        _shade(c, "F2F2F2")
        c.paragraphs[0].text = ""
        lines = text.rstrip("\n").split("\n")
        for k, line in enumerate(lines):
            para = c.paragraphs[0] if k == 0 else c.add_paragraph()
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            para.paragraph_format.first_line_indent = Cm(0)
            run = para.add_run(line if line else " ")
            _set_font(run, size=size, name="Courier New")
        self.doc.add_paragraph()

    # ------------------------------------------------------------ figuras
    def _fig_head(self, title):
        self.fig += 1
        p1 = self.p(f"**Figura {self.fig}**", indent=False, keep=True)
        p2 = self.p(f"*{title}*", indent=False, keep=True)
        return self.fig

    def _note(self, note):
        para = self.doc.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        self._runs(para, "*Nota.* " + note)
        self.doc.add_paragraph()

    def figure(self, rel_path, title, note, expect_num=None):
        """Inserta una imagen; si no existe, inserta el recuadro de captura pendiente."""
        num = self._fig_head(title)
        if expect_num is not None and expect_num != num:
            print(f"  ⚠ La figura '{rel_path}' quedó como Figura {num} (el nombre sugiere {expect_num}).")
        path = self.base / rel_path
        if path.exists():
            img = DocxImage.from_file(str(path))
            w_in = img.px_width / (img.horz_dpi or 96)
            h_in = img.px_height / (img.vert_dpi or 96)
            ratio = h_in / w_in
            width = min(TEXT_WIDTH_IN, MAX_IMG_H_IN / ratio)
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            _keep_with_next(para)
            para.add_run().add_picture(str(path), width=Inches(width))
        else:
            self.missing.append(rel_path)
            self.pending(Path(rel_path).name, note.split("|")[0] if "|" in note else title)
        self._note(note.split("|")[-1].strip() if "|" in note else note)
        return num

    def pending(self, name, what):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, {e: (12, "7F7F7F") for e in ("top", "bottom", "left", "right")})
        trPr = t.rows[0]._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        c = t.rows[0].cells[0]
        _shade(c, "D9D9D9")
        c.paragraphs[0].text = ""
        for k, (txt, b) in enumerate([("", False), (f"📸 PENDIENTE: {name}", True), (what, False), ("", False)]):
            para = c.paragraphs[0] if k == 0 else c.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            run = para.add_run(txt)
            _set_font(run, size=12 if b else 11)
            run.bold = b

    # ------------------------------------------------------------ tablas
    def table(self, title, headers, rows, widths=None, note=None, size=10):
        self.tab += 1
        self.p(f"**Tabla {self.tab}**", indent=False, keep=True)
        self.p(f"*{title}*", indent=False, keep=True)
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t, {"top": (8, "000000"), "bottom": (8, "000000"), "insideH": (2, "BFBFBF")})
        hdr = t.rows[0]
        _repeat_header(hdr)
        for k, htxt in enumerate(headers):
            c = hdr.cells[k]
            c.paragraphs[0].text = ""
            self._cellp(c.paragraphs[0], htxt, size, bold=True)
            _cell_bottom_border(c, 8)
        for row in rows:
            cells = t.add_row().cells
            for k, txt in enumerate(row):
                cells[k].paragraphs[0].text = ""
                parts = str(txt).split("\n")
                for j, part in enumerate(parts):
                    para = cells[k].paragraphs[0] if j == 0 else cells[k].add_paragraph()
                    self._cellp(para, part, size)
        if widths:
            for row in t.rows:
                for k, w in enumerate(widths):
                    row.cells[k].width = Inches(w)
        if note:
            self._note(note)
        else:
            self.doc.add_paragraph()
        return self.tab

    def _cellp(self, para, text, size, bold=False):
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        para.paragraph_format.space_after = Pt(2)
        self._runs(para, text, size=size, bold=bold)

    # ------------------------------------------------------------ especiales
    def toc(self):
        para = self.doc.add_paragraph()
        _field(para, 'TOC \\o "1-2" \\h \\z \\u')

    def reference(self, text):
        para = self.doc.add_paragraph()
        pf = para.paragraph_format
        pf.left_indent = INDENT
        pf.first_line_indent = -INDENT
        self._runs(para, text)

    def save(self, path):
        self.doc.save(path)
